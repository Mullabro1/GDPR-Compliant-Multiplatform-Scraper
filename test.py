import os
import json
import shutil
from collections import defaultdict
from math import ceil
from docx import Document
from docx.shared import Inches, RGBColor

# Base directories
script_directory = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.join(script_directory, 'test3')
DOCS_DIR = os.path.join(script_directory, 'docs3')
VIDEO_DIR = os.path.join(script_directory, 'video2')

# Ensure output dirs exist
os.makedirs(DOCS_DIR, exist_ok=True)
os.makedirs(VIDEO_DIR, exist_ok=True)

# Detect and copy only folders with videos
video_presence = {}
for folder_name in os.listdir(ROOT_DIR):
    folder_path = os.path.join(ROOT_DIR, folder_name)
    if not os.path.isdir(folder_path):
        continue
    # Check for MP4 files
    has_mp4 = any(f.lower().endswith('.mp4') for f in os.listdir(folder_path))
    video_presence[folder_name] = has_mp4
    # Only copy if MP4s exist
    if has_mp4:
        target_video_subdir = os.path.join(VIDEO_DIR, folder_name)
        shutil.copytree(folder_path, target_video_subdir, dirs_exist_ok=True)

# Collect entries per collection
collections = defaultdict(list)
for folder_name in os.listdir(ROOT_DIR):
    folder_path = os.path.join(ROOT_DIR, folder_name)
    if not os.path.isdir(folder_path):
        continue
    coll_path = os.path.join(folder_path, 'collectioninfo.json')
    data_path = os.path.join(folder_path, 'data.json')
    try:
        with open(coll_path, 'r', encoding='utf-8') as f:
            coll_info = json.load(f)
        with open(data_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
    except FileNotFoundError:
        print(f"Skipping {folder_name}: missing JSON")
        continue

    images = [os.path.join(folder_path, f) for f in os.listdir(folder_path)
              if f.lower().endswith(('.jpg', '.jpeg', '.png'))]

    entry = {
        'folder_name': folder_name,
        'url': data.get('url', ''),
        'user_posted': data.get('user_posted', ''),
        'description': data.get('description', ''),
        'hashtags': data.get('hashtags', []),
        'shortcode': data.get('shortcode', folder_name),
        'images': images,
        'has_videos': video_presence.get(folder_name, False)
    }
    for coll_name in coll_info.get('collection', []):
        collections[coll_name].append(entry)

# Generate one DOCX per collection with grouped images
for coll_name, entries in collections.items():
    safe_name = coll_name.replace(' ', '_').replace('/', '_')
    doc_path = os.path.join(DOCS_DIR, f"{safe_name}.docx")
    doc = Document()
    doc.add_heading(coll_name, level=1)

    for i, entry in enumerate(entries):
        doc.add_heading(f"Entry: {entry['shortcode']}", level=2)
        # Folder
        doc.add_paragraph(f"Folder: {entry['folder_name']}")
        # URL in blue
        p_url = doc.add_paragraph()
        p_url.add_run("URL: ")
        run_url = p_url.add_run(entry['url'])
        run_url.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)
        # User Posted
        doc.add_paragraph(f"User Posted: {entry['user_posted']}")
        # Description
        doc.add_paragraph(f"Description: {entry['description']}")
        # Hashtags
        if entry['hashtags']:
            doc.add_paragraph("Hashtags: " + ', '.join(entry['hashtags']))
        # Shortcode
        doc.add_paragraph(f"Shortcode: {entry['shortcode']}")

        # Images in 2-column table
        if entry['images']:
            num_images = len(entry['images'])
            cols = 2
            rows = ceil(num_images / cols)
            table = doc.add_table(rows=rows, cols=cols)
            table.autofit = True
            for idx, img_path in enumerate(entry['images']):
                r, c = divmod(idx, cols)
                cell = table.rows[r].cells[c]
                paragraph = cell.paragraphs[0]
                run = paragraph.add_run()
                try:
                    run.add_picture(img_path, width=Inches(3))
                except Exception as e:
                    print(f"Failed to add image {img_path}: {e}")

        # Bold, colored video note
        if entry['has_videos']:
            p_vid = doc.add_paragraph()
            run_vid = p_vid.add_run("Has MP4 but is saved for further analysis.")
            run_vid.bold = True
            run_vid.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

        # Page break between entries
        if i < len(entries) - 1:
            doc.add_page_break()

    doc.save(doc_path)
    print(f"Created: {doc_path}")
