import os
import io
import json
from collections import defaultdict
from math import ceil
from docx import Document
from docx.shared import Inches, RGBColor
from PIL import Image  # pip install --upgrade pillow

# ---- CONFIG ----
BATCH_SIZE = 50  # posts per .docx

# Base directories
script_directory = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.join(script_directory, 'test3')
DOCS_DIR = os.path.join(script_directory, 'docs3')
os.makedirs(DOCS_DIR, exist_ok=True)

# 1) Detect which folders have MP4s
video_presence = {}
for folder_name in os.listdir(ROOT_DIR):
    folder_path = os.path.join(ROOT_DIR, folder_name)
    if os.path.isdir(folder_path):
        video_presence[folder_name] = any(
            fname.lower().endswith('.mp4')
            for fname in os.listdir(folder_path)
        )

# 2) Build collections → list of entries
collections = defaultdict(list)
for folder_name in os.listdir(ROOT_DIR):
    folder_path = os.path.join(ROOT_DIR, folder_name)
    if not os.path.isdir(folder_path):
        continue

    # load JSON metadata
    try:
        with open(os.path.join(folder_path, 'collectioninfo.json'), 'r', encoding='utf-8') as f:
            coll_info = json.load(f)
        with open(os.path.join(folder_path, 'data.json'), 'r', encoding='utf-8') as f:
            data = json.load(f)
    except FileNotFoundError:
        print(f"[SKIP] {folder_name}: missing JSON")
        continue

    # gather only frame*.jpg/.jpeg images, sorted by frame number
    images = sorted(
        [
            os.path.join(folder_path, fname)
            for fname in os.listdir(folder_path)
            if fname.lower().startswith('frame') and fname.lower().endswith(('.jpg', '.jpeg'))
        ],
        key=lambda p: int(os.path.splitext(os.path.basename(p))[0][5:])
    )

    entry = {
        'folder_name': folder_name,
        'url': data.get('url', ''),
        'user_posted': data.get('user_posted', ''),
        'description': data.get('description', ''),
        'hashtags': data.get('hashtags', []),
        'shortcode': data.get('shortcode', folder_name),
        'images': images,
        'has_videos': video_presence.get(folder_name, False)
    }

    for coll_name in coll_info.get('collection', []):
        collections[coll_name].append(entry)

# 3) For each collection, split into batches and generate docs
for coll_name, entries in collections.items():
    safe_name = coll_name.replace(' ', '_').replace('/', '_')
    total_posts = len(entries)
    total_batches = ceil(total_posts / BATCH_SIZE)

    print(f"\n=== Collection: {coll_name!r} — {total_posts} posts, {total_batches} batch(es) ===")

    for batch_idx in range(total_batches):
        start = batch_idx * BATCH_SIZE
        end = start + BATCH_SIZE
        batch = entries[start:end]
        part_num = batch_idx + 1
        out_name = f"{safe_name}_{part_num}.docx"
        out_path = os.path.join(DOCS_DIR, out_name)

        print(f"[Batch {part_num}/{total_batches}] Processing posts {start+1}–{min(end, total_posts)} → {out_name}")

        doc = Document()
        doc.add_heading(coll_name, level=1)

        for i, entry in enumerate(batch, start=1):
            print(f"  • Entry {i}/{len(batch)}: {entry['shortcode']}")

            doc.add_heading(f"Entry: {entry['shortcode']}", level=2)
            doc.add_paragraph(f"Folder: {entry['folder_name']}")

            # URL in blue
            p_url = doc.add_paragraph()
            p_url.add_run("URL: ")
            run_url = p_url.add_run(entry['url'])
            run_url.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)

            doc.add_paragraph(f"User Posted: {entry['user_posted']}")
            doc.add_paragraph(f"Description: {entry['description']}")
            if entry['hashtags']:
                doc.add_paragraph("Hashtags: " + ', '.join(entry['hashtags']))
            doc.add_paragraph(f"Shortcode: {entry['shortcode']}")

            # 2-column image table, with Pillow normalization
            if entry['images']:
                num_images = len(entry['images'])
                cols = 2
                rows = ceil(num_images / cols)
                table = doc.add_table(rows=rows, cols=cols)
                table.autofit = True

                for idx, img_path in enumerate(entry['images']):
                    r, c = divmod(idx, cols)
                    cell = table.rows[r].cells[c]
                    run = cell.paragraphs[0].add_run()
                    try:
                        img = Image.open(img_path)
                        if img.mode in ("RGBA", "P", "CMYK"):
                            img = img.convert("RGB")
                        buf = io.BytesIO()
                        img.save(buf, format="JPEG")
                        buf.seek(0)
                        run.add_picture(buf, width=Inches(3))
                    except Exception as e:
                        print(f"    [WARN] Failed to embed {os.path.basename(img_path)}: {e}")

            # Video note
            if entry['has_videos']:
                p_vid = doc.add_paragraph()
                run_vid = p_vid.add_run("Has MP4 but is saved for further analysis.")
                run_vid.bold = True
                run_vid.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

            # page break except after last entry
            if i < len(batch):
                doc.add_page_break()

        # save batch doc
        doc.save(out_path)
        print(f"[OK] Saved {out_name}")

print("\nAll batches complete.")  
